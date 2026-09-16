#!/usr/bin/env python3
"""스마트스토어/쿠팡 → 사방넷 발주등록 변환"""
import argparse,json,os,re,sys
import pandas as pd
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,Alignment,Border,Side,PatternFill
SD=os.path.dirname(os.path.abspath(__file__))
MF=os.path.join(SD,"product_mapping.json")
H=["상품고유코드","판매상품명","수량","배송방식","주문자 이름","받는분 이름","전화번호1","전화번호2","우편번호","주소1","주소2","배송메세지","주문번호","관리메모1","관리메모2","관리메모3","관리메모4","관리메모5","상품별 메모1","상품별 메모2","상품별 메모3","발주 타입","출고희망일"]
RQ={"상품고유코드","판매상품명","수량","배송방식","받는분 이름","전화번호1","우편번호","주소1"}
def ld_map():
    import zlib,base64
    if os.path.exists(MF):
        try:
            with open(MF,"r",encoding="utf-8") as f: return json.load(f)
        except: pass
    try:
        sys.path.insert(0,SD); from mapping_data import D
        return json.loads(zlib.decompress(base64.b85decode(D)).decode("utf-8"))
    except: pass
    return {"상품번호별_매핑":{},"쿠팡_매핑":{},"코드별_상품명":{}}
def detect(p):
    fn=os.path.basename(p).lower()
    if "스마트스토어" in fn or "smartstore" in fn: return "smartstore"
    if "deliverylist" in fn or "delivery" in fn: return "coupang"
    wb=load_workbook(p,read_only=True); sn=[s.lower() for s in wb.sheetnames]; wb.close()
    if any("발주발송관리" in s for s in sn): return "smartstore"
    if any("delivery" in s for s in sn): return "coupang"
    return "smartstore"
def parse_opt(o):
    c,sz,st,eq=None,None,None,None; opt=str(o)
    for p in opt.split(" / "):
        p=p.strip()
        if p.startswith("색상:") or p.startswith("컬러:"): c=p.split(":",1)[1].strip()
        elif p.startswith("사이즈:"): sz=p.replace("사이즈:","").strip()
        elif p.startswith("수량:") and not re.search(r"\d+개",p): sz=p.split(":",1)[1].strip()
        elif p.startswith("세트"): st=p
    m=re.search(r"할인이벤트:\s*(\d+)개",opt)
    if m: eq=int(m.group(1))
    return c,sz,st,eq
def set_mul(st):
    if not st: return 1
    if "2+2" in st: return 2
    if "3+3" in st: return 3
    return 1
def get_code(mp,pno,color,size):
    pno=str(pno); pm=mp.get("상품번호별_매핑",{})
    if pno not in pm: return "???"
    i=pm[pno]
    if "_default" in i.get("매핑",{}): return i["매핑"]["_default"]
    if i.get("색상필요") and color and size: return i["매핑"].get(f"{color}_{size}","???")
    elif size: return i["매핑"].get(size,"???")
    elif len(i.get("매핑",{}))==1: return list(i["매핑"].values())[0]
    return "???"
def cq(o):
    opt=str(o); m=re.search(r"(\d+)개",opt)
    if m: return int(m.group(1))
    m=re.search(r"(\d+)박스",opt)
    if m: return int(m.group(1))
    return 1
def cc(mp,pn,on):
    cm=mp.get("쿠팡_매핑",{}); p,o=str(pn).strip(),str(on).strip()
    if f"{p}|{o}" in cm: return cm[f"{p}|{o}"]
    if f"{p}|*" in cm: return cm[f"{p}|*"]
    return "???"
def cl(v): return "" if str(v) in ("nan","None") else str(v)
def fz(v):
    s=cl(v).replace(".0","")
    if s and s.isdigit(): s=s.zfill(5)
    return s
def ac(pc):
    if pc.startswith("N"): return "뉴트리정"
    elif pc.startswith("E"): return "이더컴퍼니"
    return "unknown"
def mr(pc,pn,q,rcv,ph,zc,addr,msg,ono):
    return {"상품고유코드":pc,"판매상품명":pn,"수량":q,"배송방식":"택배","주문자 이름":"","받는분 이름":cl(rcv),"전화번호1":cl(ph),"전화번호2":"","우편번호":fz(zc),"주소1":cl(addr),"주소2":"","배송메세지":cl(msg),"주문번호":cl(ono),"관리메모1":"","관리메모2":"","관리메모3":"","관리메모4":"","관리메모5":"","상품별 메모1":"","상품별 메모2":"","상품별 메모3":"","발주 타입":"","출고희망일":None}
def decrypt_if_needed(ip):
    try:
        with open(ip,"rb") as f: head=f.read(8)
    except: return ip
    if head[:4]==b"\xd0\xcf\x11\xe0":
        try:
            import msoffcrypto, tempfile
            tmp=tempfile.NamedTemporaryFile(suffix=".xlsx",delete=False).name
            with open(ip,"rb") as f:
                of=msoffcrypto.OfficeFile(f)
                if of.is_encrypted():
                    of.load_key(password="1234")
                    with open(tmp,"wb") as o: of.decrypt(o)
                    print(f"  → 암호 해제 완료 (임시: {tmp})")
                    return tmp
        except Exception as e:
            print(f"  ⚠️  암호 해제 실패: {e}")
    return ip
def conv_ss(ip,mp):
    ip=decrypt_if_needed(ip)
    df=pd.read_excel(ip,header=1); ar={}; um=[]
    for i,r in df.iterrows():
        c,sz,st,eq=parse_opt(r.get("옵션정보",""))
        pno=str(r.get("상품번호","")); pc=get_code(mp,pno,c,sz); pn=mp.get("코드별_상품명",{}).get(pc,"")
        bq=int(r.get("수량",1)); q=eq*bq if eq else bq*set_mul(st)
        if pc=="???": um.append({"행":i+3,"상품번호":pno,"상품명":str(r.get("상품명","")),"옵션정보":str(r.get("옵션정보",""))})
        ar.setdefault(ac(pc),[]).append(mr(pc,pn,q,r.get("수취인명",""),r.get("수취인연락처1",""),r.get("우편번호",""),r.get("통합배송지",""),r.get("배송메세지",""),r.get("상품주문번호","")))
    return ar,um
def conv_cpg(ip,mp):
    df=pd.read_excel(ip,header=0); df=df.dropna(subset=["등록상품명"]); ar={}; um=[]
    for i,r in df.iterrows():
        pn,on=str(r.get("등록상품명","")),str(r.get("등록옵션명",""))
        pc=cc(mp,pn,on); sn=mp.get("코드별_상품명",{}).get(pc,"")
        bq=r.get("구매수(수량)",1); q=cq(on)*int(1 if pd.isna(bq) else bq)
        if pc=="???": um.append({"행":i+2,"등록상품명":pn,"등록옵션명":on})
        ar.setdefault(ac(pc),[]).append(mr(pc,sn,q,r.get("수취인이름",""),r.get("수취인전화번호",""),r.get("우편번호",""),r.get("수취인 주소",""),r.get("배송메세지",""),r.get("주문번호","")))
    return ar,um
def wr_out(mg,op,sp):
    of={}
    if sp:
        b,e=os.path.splitext(op)
        for a,rows in mg.items():
            if not rows: continue
            p=f"{b}_{a}{e}"; wr_xl(p,rows); of[a]={"path":p,"count":len(rows)}
    else:
        al=[]; [al.extend(r) for r in mg.values()]; wr_xl(op,al); of["전체"]={"path":op,"count":len(al)}
    return of
def wr_xl(op,rows):
    rows=sorted(rows,key=lambda x:str(x.get("판매상품명","")))
    wb=Workbook(); ws=wb.active; ws.title="발주등록_sample"
    rf=Font(name="Arial",bold=True,color="FF0000",size=10); bf=Font(name="Arial",bold=True,size=10)
    df=Font(name="Arial",size=10); hf=PatternFill("solid",fgColor="F2F2F2")
    bd=Border(left=Side(style="thin"),right=Side(style="thin"),top=Side(style="thin"),bottom=Side(style="thin"))
    for ci,h in enumerate(H,1):
        cell=ws.cell(row=1,column=ci,value=h); cell.font=rf if h in RQ else bf; cell.fill=hf; cell.alignment=Alignment(horizontal="center",vertical="center"); cell.border=bd
    text_cols={"주문번호","전화번호1","전화번호2","우편번호","상품고유코드"}
    for ri,r in enumerate(rows,2):
        for ci,h in enumerate(H,1):
            v=r[h]
            if h in text_cols and v not in (None,""): v=str(v)
            cell=ws.cell(row=ri,column=ci,value=v); cell.font=df; cell.border=bd
            if h in text_cols: cell.number_format="@"
    for c,w in {"A":14,"B":45,"C":6,"D":8,"E":12,"F":12,"G":16,"H":16,"I":10,"J":55,"K":20,"L":25,"M":20}.items(): ws.column_dimensions[c].width=w
    wb.save(op)
def main():
    pa=argparse.ArgumentParser(); pa.add_argument("--input",required=True,nargs="+"); pa.add_argument("--output",required=True)
    pa.add_argument("--source",choices=["smartstore","coupang"]); pa.add_argument("--no-split",action="store_true")
    a=pa.parse_args(); mp=ld_map(); aa=[]; au=[]
    for inp in a.input:
        src=a.source or detect(inp); print(f"입력 파일: {os.path.basename(inp)} (형식: {src})")
        if src=="smartstore": ar,um=conv_ss(inp,mp)
        elif src=="coupang": ar,um=conv_cpg(inp,mp)
        else: print(f"알 수 없는 소스: {src}"); sys.exit(1)
        aa.append(ar); au.extend(um)
    mg={}
    for ar in aa:
        for a2,rows in ar.items(): mg.setdefault(a2,[]).extend(rows)
    of=wr_out(mg,a.output,sp=not a.no_split); tot=sum(len(r) for r in mg.values())
    print(f"\n변환 완료: 총 {tot}건")
    for a2,info in of.items(): print(f"  [{a2}] {info['count']}건 → {info['path']}")
    if au:
        print(f"\n⚠️  매핑 실패: {len(au)}건")
        for u in au:
            if "상품번호" in u: print(f"  행 {u['행']}: [{u['상품번호']}] {u['상품명'][:30]}... | {u['옵션정보'][:40]}...")
            else: print(f"  행 {u['행']}: {u['등록상품명'][:40]} | {u['등록옵션명'][:30]}")
    else: print("모든 상품 매핑 성공!")
if __name__=="__main__": main()
