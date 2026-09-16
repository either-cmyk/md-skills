#!/usr/bin/env python3
"""
빈박스 운송장 자동 매칭 스크립트

쿠팡 DeliveryList의 주문번호와 3PL 체험단 송장번호를 매칭하여
운송장번호 복사용 엑셀 파일을 생성한다.

3PL 송장 파일은 두 가지 방법으로 지정할 수 있다:
  1. --kakao-dir 옵션으로 카카오톡 받은 파일 폴더를 지정하면
     오늘 날짜의 3pl 파일을 자동으로 찾아서 사용 (권장)
  2. --tracking 옵션으로 파일을 직접 지정

--search-days N 옵션을 사용하면 오늘부터 N일 전까지의 3PL 파일도 함께 탐색한다.
미매칭 건이 있을 때 이전 날짜 파일에서 추가 매칭을 시도하는 데 유용하다.

사용법:
    # 방법 1: 카카오톡 폴더에서 3PL 자동 탐색 (권장)
    python3 process_delivery.py \
        --delivery "DeliveryList(2026-04-07)_(0).xlsx" \
        --kakao-dir "/path/to/카카오톡 받은 파일" \
        --output-dir "~/Downloads"

    # 방법 2: 이전 날짜까지 포함하여 탐색 (미매칭 재처리용)
    python3 process_delivery.py \
        --delivery "DeliveryList(2026-04-07)_(0).xlsx" \
        --kakao-dir "/path/to/카카오톡 받은 파일" \
        --search-days 7 \
        --output-dir "~/Downloads"

    # 방법 3: 3PL 파일 직접 지정
    python3 process_delivery.py \
        --delivery "DeliveryList(2026-04-07)_(0).xlsx" \
        --tracking "3pl파일1.xlsx" "3pl파일2.xlsx" \
        --output-dir "~/Downloads"
"""

import argparse
import os
import sys
from datetime import datetime, timedelta

try:
    import openpyxl
except ImportError:
    os.system("pip install openpyxl --break-system-packages -q")
    import openpyxl

try:
    import xlsxwriter
except ImportError:
    os.system("pip install xlsxwriter --break-system-packages -q")
    import xlsxwriter


def find_3pl_files(kakao_dir, search_days=0):
    """카카오톡 받은 파일 폴더에서 3PL 파일을 자동으로 찾는다.

    파일명 패턴: '3pl C 체험단 [상품명] MM DD_[숫자].xlsx'

    search_days=0이면 오늘 날짜만, N이면 오늘부터 N일 전까지 탐색한다.
    """
    today = datetime.now()
    date_patterns = []
    for i in range(search_days + 1):
        d = today - timedelta(days=i)
        date_patterns.append(d.strftime("%m %d"))   # e.g. "04 07"
        date_patterns.append(d.strftime("%m_%d"))    # e.g. "04_07"

    found_files = []
    if not os.path.isdir(kakao_dir):
        print(f"  [오류] 카카오톡 받은 파일 폴더를 찾을 수 없습니다: {kakao_dir}")
        return found_files

    for fname in os.listdir(kakao_dir):
        fname_lower = fname.lower()
        if "3pl" not in fname_lower:
            continue
        if not fname.endswith(".xlsx"):
            continue
        if any(pat in fname for pat in date_patterns):
            full_path = os.path.join(kakao_dir, fname)
            found_files.append(full_path)

    return sorted(found_files)


def extract_tracking_mapping(tracking_files):
    """3PL 체험단 파일들에서 주문번호 → 송장번호 매핑을 추출한다."""
    mapping = {}
    duplicates = []

    for fpath in tracking_files:
        wb = openpyxl.load_workbook(fpath, data_only=True)
        ws = wb.active

        for row in range(2, ws.max_row + 1):
            order_raw = ws.cell(row=row, column=1).value
            cj_tracking = ws.cell(row=row, column=11).value
            lotte_tracking = ws.cell(row=row, column=12).value

            if not order_raw:
                continue

            tracking_raw = cj_tracking if cj_tracking else lotte_tracking
            if not tracking_raw:
                continue

            try:
                order_int = int(str(order_raw).strip().replace(".0", ""))
            except (ValueError, TypeError):
                continue

            tracking_str = str(tracking_raw).strip().replace("-", "").replace(".0", "")
            try:
                tracking_int = int(tracking_str)
            except (ValueError, TypeError):
                continue

            if order_int in mapping and mapping[order_int] != tracking_int:
                duplicates.append((order_int, mapping[order_int], tracking_int))

            mapping[order_int] = tracking_int

        wb.close()

    return mapping, duplicates


def process_delivery_list(delivery_path, mapping, output_dir):
    """DeliveryList에서 주문번호를 읽고, 매칭된 운송장번호를 복사용 엑셀로 출력한다."""
    wb = openpyxl.load_workbook(delivery_path, data_only=True)
    ws = wb.active

    matched = 0
    unmatched = 0
    tracking_list = []

    for row in range(2, ws.max_row + 1):
        order_val = ws.cell(row=row, column=3).value
        if not order_val:
            continue
        try:
            order_int = int(str(order_val).strip().replace(".0", ""))
        except (ValueError, TypeError):
            tracking_list.append(None)
            continue

        if order_int in mapping:
            tracking_list.append(mapping[order_int])
            matched += 1
        else:
            tracking_list.append(None)
            unmatched += 1

    wb.close()

    # xlsxwriter로 복사용 엑셀 생성
    os.makedirs(output_dir, exist_ok=True)
    xlsx_output = os.path.join(output_dir, "운송장번호_복사용.xlsx")
    wb_out = xlsxwriter.Workbook(xlsx_output)
    ws_out = wb_out.add_worksheet("운송장번호")
    num_fmt = wb_out.add_format({'num_format': '0'})

    for i, val in enumerate(tracking_list):
        if val is not None:
            ws_out.write_number(i, 0, val, num_fmt)

    wb_out.close()

    # 결과 출력
    print(f"\n{'='*60}")
    print(f"처리 결과:")
    print(f"  매칭 성공: {matched}건")
    print(f"  매칭 실패(미입력): {unmatched}건")
    print(f"{'='*60}")
    print(f"\n출력 파일: {xlsx_output}")
    print(f"\n사용법: 다운로드 폴더에서 운송장번호_복사용.xlsx 열기")
    print(f"      → A열 드래그 → 복사(Ctrl+C)")
    print(f"      → DeliveryList 운송장번호(E열) 첫 칸에 붙여넣기(Ctrl+V)")

    return {
        "matched": matched,
        "unmatched": unmatched,
        "xlsx_output": xlsx_output,
        "tracking_list": tracking_list,
    }


def main():
    parser = argparse.ArgumentParser(description="빈박스 운송장 자동 매칭")
    parser.add_argument("--delivery", required=True, help="DeliveryList 엑셀 파일 경로")
    parser.add_argument("--tracking", nargs="+", help="3PL 체험단 송장 엑셀 파일 경로들 (직접 지정)")
    parser.add_argument("--kakao-dir", help="카카오톡 받은 파일 폴더 경로 (3PL 자동 탐색)")
    parser.add_argument("--search-days", type=int, default=0, help="탐색할 이전 날짜 수 (기본: 0=오늘만, 7=일주일)")
    parser.add_argument("--output-dir", required=True, help="출력 폴더 경로 (다운로드 폴더)")
    args = parser.parse_args()

    # 3PL 파일 결정: --tracking이 있으면 직접 지정, 없으면 --kakao-dir에서 자동 탐색
    tracking_files = []

    if args.tracking:
        tracking_files = args.tracking
        print(f"3PL 송장 파일 {len(tracking_files)}개 (직접 지정):")
        for f in tracking_files:
            print(f"  - {os.path.basename(f)}")
    elif args.kakao_dir:
        if args.search_days > 0:
            print(f"카카오톡 받은 파일 폴더에서 최근 {args.search_days + 1}일간 3PL 파일 탐색 중...")
        else:
            print(f"카카오톡 받은 파일 폴더에서 오늘 날짜 3PL 파일 탐색 중...")
        print(f"  폴더: {args.kakao_dir}")
        tracking_files = find_3pl_files(args.kakao_dir, args.search_days)
        if tracking_files:
            print(f"  발견된 3PL 파일 {len(tracking_files)}개:")
            for f in tracking_files:
                print(f"  - {os.path.basename(f)}")
        else:
            today = datetime.now()
            date_str = today.strftime("%m %d")
            print(f"  [오류] 3PL 파일을 찾을 수 없습니다.")
            print(f"  카카오톡에서 3PL 파일을 다운로드했는지 확인해주세요.")
            sys.exit(1)
    else:
        print("[오류] --tracking 또는 --kakao-dir 중 하나를 지정해야 합니다.")
        sys.exit(1)

    print(f"\n3PL 송장 파일에서 매핑 추출 중...")
    mapping, duplicates = extract_tracking_mapping(tracking_files)
    print(f"  총 매핑: {len(mapping)}건")

    if duplicates:
        print(f"\n  [주의] 중복 주문번호 {len(duplicates)}건:")
        for order, old, new in duplicates:
            print(f"    주문 {order}: {old} → {new} (후자 사용)")

    print(f"\nDeliveryList 처리 중...")
    process_delivery_list(args.delivery, mapping, args.output_dir)


if __name__ == "__main__":
    main()
