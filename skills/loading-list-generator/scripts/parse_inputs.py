#!/usr/bin/env python3
"""동봉문서 PDF + CN 출고내역 엑셀 → 적재리스트 JSON 생성.

규칙 (2026.04.21):
- 페이지 1개 = 팔레트 1개
- 박스번호 = {해당 팔레트 박스수}-{팔레트 내 박스 순번}
- 혼적박스만 박스번호 기재, 단독박스 상품은 SKU별 합산 1행 (box_no 공란)
- 팔레트 페이지에는 해당 팔레트 상품만 나열
"""
import argparse, json, re, subprocess, sys
from collections import OrderedDict, defaultdict

try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "--break-system-packages", "-q"])
    import openpyxl


def parse_dongbong_pdf(pdf_path: str) -> dict:
    """동봉문서 PDF → {업체명, 업체번호, 발주번호, 센터, 도착예정일, 팔레트수량, items[]}

    여러 레이아웃을 처리:
    - sku/바코드 같은 줄 ("62994252 /8802937071547")
    - sku와 바코드 다른 줄
    - sku 뒤에 옵션 텍스트가 붙는 경우
    """
    text = subprocess.check_output(["pdftotext", "-layout", pdf_path, "-"]).decode("utf-8")
    head = text.split("업체 보관용")[0]

    def extract(pat, src=head, group=1, default=""):
        m = re.search(pat, src)
        return m.group(group).strip() if m else default

    info = {
        "company_name": extract(r"업체명\s+(\S+)"),
        "company_code": extract(r"업체번호\s+(\S+)"),
        "order_id": extract(r"발주번호\s+(\d+)"),
        "center": extract(r"납품 센터\s+(\S+)"),
        "pallet_count": int(extract(r"팔레트수량\s+(\d+)", default="0")),
    }
    dd = extract(r"물류센터 도착예정일\s+(\d{8})")
    info["delivery_date"] = f"{dd[0:4]}-{dd[4:6]}-{dd[6:8]}" if len(dd) == 8 else ""

    body_m = re.search(r"3\.상품 정보(.+?)합계", head, re.DOTALL)
    items = []
    if not body_m:
        info["items"] = items
        return info

    body = body_m.group(1)
    lines = body.split("\n")

    # 아이템 번호 + 수량 라인 찾기:
    # "   N    (sku?)    ...   qty   qty   ..."
    # N=1~3 digit, qty = 같은 수 2번 반복
    item_line_re = re.compile(
        r"^\s*(\d{1,3})\s+(?:(\d{8})\s+)?.*?(\d+)\s+(\3)(?:\s|$)"
    )

    # 각 아이템 라인을 찾고, 블록 범위를 결정 (이전 아이템 이후 ~ 현재 아이템 라인까지)
    item_line_idx = []  # (line_idx, no, sku_or_None, qty)
    for i, line in enumerate(lines):
        m = item_line_re.match(line)
        if m:
            item_line_idx.append((i, int(m.group(1)), m.group(2), int(m.group(3))))

    # 블록 범위: 이전 아이템 라인 +1 ~ 현재 아이템 라인 + 2 (바코드가 다음 줄에 올 수도 있음)
    for idx, (line_no, no, sku_inline, qty) in enumerate(item_line_idx):
        start = item_line_idx[idx - 1][0] + 1 if idx > 0 else 0
        end = min(line_no + 3, len(lines))
        block_text = "\n".join(lines[start:end])

        # SKU: 아이템 라인에 인라인으로 있거나 블록 내 별도 8자리 숫자
        sku = sku_inline
        if not sku:
            m = re.search(r"(?<!\d)(\d{8})(?!\d)", block_text)
            if m:
                sku = m.group(1)

        # 바코드: 인라인 "sku /barcode" 혹은 별도 13~15자리 숫자
        barcode = ""
        m = re.search(r"/(\d{13,15})", block_text)
        if m:
            barcode = m.group(1)
        else:
            # 블록 내 13~15자리 숫자 (sku와 겹치지 않는)
            cands = re.findall(r"(?<!\d)(\d{13,15})(?!\d)", block_text)
            for c in cands:
                if c != sku:
                    barcode = c
                    break

        # 테이블 헤더 단어 블랙리스트
        HEADER_WORDS = {
            "상품명", "옵션", "업체발주", "제조일자관리", "제조", "수입일자",
            "상품번호", "상품", "바코드", "발주수량", "비고", "확정수량",
            "소비기한관리", "소비기한", "Box", "No",
        }

        def tokenize(s):
            # 공백/슬래시/괄호 기준으로 한글·영문 토큰 추출
            return re.findall(r"[가-힣A-Za-z0-9~]+", s)

        def is_header_only_line(s):
            toks = re.findall(r"[가-힣A-Za-z]+", s)
            if not toks: return False
            return all(t in HEADER_WORDS for t in toks)

        def clean_name_line(s):
            """한 줄에서 상품명 부분만 추출. 숫자/sku/헤더 단어 제거."""
            # sku/바코드 패턴 제거
            s = re.sub(r"\d{8}\s*/?\s*\d{13,15}", " ", s)
            s = re.sub(r"(?<!\d)\d{13,15}(?!\d)", " ", s)
            s = re.sub(r"(?<!\d)\d{8}(?!\d)", " ", s)
            # 토큰 분리 (공백 기준) 후 헤더 토큰 제거 — 한 토큰에 붙어있는 경우 하위 분리
            out_tokens = []
            for tok in re.split(r"\s+", s.strip()):
                if not tok: continue
                # 토큰 내부에도 헤더 단어가 슬래시/괄호로 붙을 수 있음
                sub = re.split(r"[/()]+", tok)
                sub_clean = [p for p in sub if p and p not in HEADER_WORDS]
                if sub_clean and all(p in HEADER_WORDS for p in sub):
                    # 전체가 헤더 단어면 skip
                    continue
                # 단독 "N"/"Y"/"-" 토큰 제거
                if re.fullmatch(r"[NY\-]", tok): continue
                # 괄호·슬래시로 엮인 부분에서 헤더 아닌 조각만 다시 조합
                kept_pieces = [p for p in sub if p and p not in HEADER_WORDS]
                if kept_pieces:
                    out_tokens.append("/".join(kept_pieces) if len(kept_pieces) > 1 else kept_pieces[0])
            return " ".join(out_tokens).strip()

        name_parts = []
        for bl in lines[start:line_no]:
            s = bl.strip()
            if not s: continue
            if re.match(r"^[\d\s/\-]+$", s): continue
            if is_header_only_line(s): continue
            cleaned = clean_name_line(s)
            if not cleaned: continue
            if re.search(r"[가-힣]", cleaned):
                name_parts.append(cleaned)
        name = " ".join(name_parts)
        name = re.sub(r"\s+", " ", name).strip()
        # 상품명 정리: "/"로 상품명/옵션 구분하되 앞뒤 정돈
        name = re.sub(r"\s*/\s*", " / ", name)

        items.append({
            "no": no,
            "sku": sku or "",
            "barcode": barcode,
            "name": name,
            "qty": qty,
        })

    info["items"] = items
    return info

def parse_cn_excel(xlsx_path: str):
    """CN 출고내역 엑셀 → (팔레트별 박스 구조, 엑셀 총박스수)"""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.worksheets[0]
    
    row1 = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0][0] or ""
    m = re.search(r"박스수량[:：]\s*(\d+)", str(row1))
    excel_total_boxes = int(m.group(1)) if m else None
    
    # 헤더 감지: 2행에 "트레이 번호" 있는지
    row2 = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
    headers = [str(h or "").strip() for h in row2]
    # 컬럼 인덱스
    def idx(*names):
        for n in names:
            if n in headers: return headers.index(n)
        return -1
    
    col_tray = idx("트레이 번호")
    col_box = idx("상자 번호")
    col_sku = idx("SKU")
    col_label = idx("라벨명")
    col_qty = idx("출고수량")
    
    pallets = OrderedDict()
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or all(v is None for v in row): continue
        tray = row[col_tray] if col_tray >= 0 else None
        box = row[col_box] if col_box >= 0 else None
        sku = row[col_sku] if col_sku >= 0 else None
        label = row[col_label] if col_label >= 0 else ""
        qty = row[col_qty] if col_qty >= 0 else 0
        if not box or not sku: continue
        box = str(box); sku = str(sku); q = float(qty or 0)
        if tray not in pallets:
            pallets[tray] = OrderedDict()  # box -> [(sku, qty)]
        if box not in pallets[tray]:
            pallets[tray][box] = []
        pallets[tray][box].append((sku, q))
    
    return pallets, excel_total_boxes


def build_loading_list(dongbong: dict, pallets: OrderedDict) -> dict:
    """동봉 정보 + 팔레트별 박스 데이터 → 적재리스트 JSON"""
    # 바코드 → (no, sku, name)
    bc_to_item = {}
    for it in dongbong["items"]:
        bc = it.get("barcode", "")
        if bc:
            bc_to_item[bc] = (it["no"], it["sku"], it["name"])
    
    result_pallets = []
    warnings = []
    
    for seq, (tray, boxes) in enumerate(pallets.items(), 1):
        box_count = len(boxes)
        single_sum = defaultdict(float)  # bc -> qty (단독박스 합산)
        mix_rows = []  # (bc, box_seq, qty)
        
        for box_seq, (box_id, items) in enumerate(boxes.items(), 1):
            distinct_bcs = set(bc for bc,_ in items)
            if len(distinct_bcs) >= 2:
                for bc, q in items:
                    if bc in bc_to_item:
                        mix_rows.append((bc, box_seq, q))
            else:
                for bc, q in items:
                    if bc in bc_to_item:
                        single_sum[bc] += q
        
        rows = []
        row_no = 0
        # 단독: 동봉 순번으로 정렬
        for bc, total_q in sorted(single_sum.items(), key=lambda x: bc_to_item[x[0]][0]):
            no, sku, name = bc_to_item[bc]
            row_no += 1
            rows.append({"no": row_no, "sku": sku, "name": name, "box_no": "", "qty": int(total_q)})
        # 혼적: 박스순 + 동봉 순번
        for bc, box_seq, q in sorted(mix_rows, key=lambda x: (x[1], bc_to_item[x[0]][0])):
            no, sku, name = bc_to_item[bc]
            row_no += 1
            rows.append({
                "no": row_no, "sku": sku, "name": name,
                "box_no": f"{box_count}-{box_seq}", "qty": int(q),
            })
        
        if rows:
            result_pallets.append({"seq": seq, "box_count": box_count, "rows": rows})
    
    # 동봉 팔레트수 vs CN 팔레트수 검증
    if dongbong["pallet_count"] and len(pallets) != dongbong["pallet_count"]:
        warnings.append(
            f"팔레트수 불일치: 동봉={dongbong['pallet_count']}, CN엑셀={len(pallets)}"
        )
    
    data = {
        "company_name": dongbong["company_name"],
        "company_code": dongbong["company_code"],
        "order_id": dongbong["order_id"],
        "center": dongbong["center"],
        "delivery_date": dongbong["delivery_date"],
        "pallet_count": dongbong["pallet_count"] or len(pallets),
        "total_boxes": sum(p["box_count"] for p in result_pallets) if result_pallets else sum(len(bs) for bs in pallets.values()),
        "pallets": result_pallets,
        "_warnings": warnings,
    }
    return data


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--dongbong-pdf", required=True)
    ap.add_argument("--cn-excel", required=True)
    ap.add_argument("--output-json", required=True)
    args = ap.parse_args()
    
    dongbong = parse_dongbong_pdf(args.dongbong_pdf)
    pallets, excel_total = parse_cn_excel(args.cn_excel)
    data = build_loading_list(dongbong, pallets)
    
    with open(args.output_json, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    
    print(json.dumps({
        "status": "ok",
        "order_id": data["order_id"],
        "pallet_count": data["pallet_count"],
        "excel_total_boxes": excel_total,
        "pallets_with_dongbong_items": len(data["pallets"]),
        "dongbong_item_count": len(dongbong["items"]),
        "warnings": data["_warnings"],
    }, ensure_ascii=False, indent=2))
