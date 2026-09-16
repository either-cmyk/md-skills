#!/usr/bin/env python3
"""CN인사이더 출고내역 → 사방넷 간편입고 등록 변환 스크립트"""

import sys
import re
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from collections import defaultdict


def load_barcode_map(product_code_path):
    """공산품 상품코드 엑셀에서 바코드→(상품코드, 출고상품명) 매핑 로드"""
    wb = openpyxl.load_workbook(product_code_path)
    ws = wb.active
    barcode_map = {}
    for r in range(2, ws.max_row + 1):
        code = ws.cell(row=r, column=3).value    # C: 상품코드
        name = ws.cell(row=r, column=4).value     # D: 출고상품명
        barcode = ws.cell(row=r, column=6).value  # F: 바코드
        if barcode and code:
            barcode_map[str(barcode).strip()] = (str(code).strip(), name)
    return barcode_map


def parse_shipment(shipment_path):
    """출고내역 엑셀 파싱 → 상품별 그룹핑된 데이터 반환"""
    wb = openpyxl.load_workbook(shipment_path)
    ws = wb.active

    # 1행 타이틀에서 S-QED 코드 추출
    title = str(ws.cell(row=1, column=1).value or '')
    sqed_match = re.search(r'(S-QED\d+)', title)
    sqed_code = sqed_match.group(1) if sqed_match else 'S-QED'

    rows = []
    for r in range(3, ws.max_row + 1):
        box_no = ws.cell(row=r, column=2).value    # B: 상자번호
        sku = ws.cell(row=r, column=6).value        # F: SKU/바코드
        label = ws.cell(row=r, column=8).value       # H: 라벨명
        qty_k = ws.cell(row=r, column=11).value      # K: 출고수량
        qty_l = ws.cell(row=r, column=12).value      # L: 세트수량

        if not label or str(label).strip() == '':
            continue
        if box_no and str(box_no).strip().upper() == 'TOTAL':
            continue

        qty = qty_l if qty_l and float(qty_l) > 0 else qty_k
        qty = int(float(qty)) if qty else 0
        sku_str = str(sku).strip() if sku else ''
        box_str = str(box_no).strip() if box_no else ''

        rows.append({
            'box': box_str,
            'sku': sku_str,
            'label': label,
            'qty': qty
        })

    # 원본 총수량 (검증용)
    total_row = ws.cell(row=ws.max_row, column=2).value
    original_total_l = ws.cell(row=ws.max_row, column=12).value
    original_total_k = ws.cell(row=ws.max_row, column=11).value

    return rows, sqed_code, original_total_l, original_total_k


def group_by_product(rows, barcode_map):
    """바코드 기준 상품 그룹핑 + 바코드 매핑"""
    product_data = defaultdict(lambda: {
        'boxes': set(), 'qty': 0, 'label': '',
        'prod_code': '', 'prod_name': '', 'barcode': ''
    })

    unmatched = []
    for row in rows:
        key = row['sku']
        if key in barcode_map:
            prod_code, prod_name = barcode_map[key]
        else:
            prod_code, prod_name = '', row['label']
            if key not in [u['sku'] for u in unmatched]:
                unmatched.append({'sku': key, 'label': row['label']})

        product_data[key]['boxes'].add(row['box'])
        product_data[key]['qty'] += row['qty']
        product_data[key]['label'] = row['label']
        product_data[key]['prod_code'] = prod_code
        product_data[key]['prod_name'] = prod_name
        product_data[key]['barcode'] = key

    # 라벨명 오름차순 정렬
    sorted_products = sorted(product_data.values(), key=lambda x: x['label'])
    return sorted_products, unmatched


def format_box_range(boxes):
    """박스번호 세트 → 범위 문자열 (S-QED-N-001~S-QED-N-003, S-QED-N-005)"""
    nums = []
    prefix = ''
    for b in boxes:
        parts = b.rsplit('-', 1)
        if len(parts) == 2:
            prefix = parts[0]
            try:
                nums.append(int(parts[1]))
            except ValueError:
                continue

    if not nums:
        return ', '.join(sorted(boxes))

    nums.sort()
    ranges = []
    i = 0
    while i < len(nums):
        start = nums[i]
        end = start
        while i + 1 < len(nums) and nums[i + 1] == end + 1:
            i += 1
            end = nums[i]
        if start == end:
            ranges.append(f"{prefix}-{start:03d}")
        else:
            ranges.append(f"{prefix}-{start:03d}~{prefix}-{end:03d}")
        i += 1
    return ', '.join(ranges)


def create_output(products, output_path):
    """사방넷 간편입고 등록 양식 엑셀 생성"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '엑셀입고'

    headers = ['박스넘버', '출고상품코드', '상품명', '수량', '바코드', '유통기한', '로케이션', '입고메모']
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, size=10, name='맑은 고딕')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    for i, p in enumerate(products, 2):
        box_range = format_box_range(p['boxes'])
        data = [box_range, p['prod_code'], p['prod_name'], p['qty'], p['barcode'], None, None, None]
        for c, val in enumerate(data, 1):
            cell = ws.cell(row=i, column=c, value=val)
            cell.font = Font(size=10, name='맑은 고딕')
            cell.border = thin_border

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 55
    ws.column_dimensions['D'].width = 8
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 12

    wb.save(output_path)
    return output_path


def main():
    if len(sys.argv) < 3:
        print("사용법: python convert.py <출고내역.xlsx> <공산품상품코드.xlsx> [출력파일.xlsx]")
        sys.exit(1)

    shipment_path = sys.argv[1]
    product_code_path = sys.argv[2]
    output_path = sys.argv[3] if len(sys.argv) > 3 else None

    # 바코드 매핑 로드
    barcode_map = load_barcode_map(product_code_path)
    print(f"공산품 바코드 매핑: {len(barcode_map)}건 로드")

    # 출고내역 파싱
    rows, sqed_code, orig_total_l, orig_total_k = parse_shipment(shipment_path)
    print(f"출고내역: {len(rows)}건 로드 (코드: {sqed_code})")

    # 상품 그룹핑
    products, unmatched = group_by_product(rows, barcode_map)

    if unmatched:
        print(f"\n⚠️ 바코드 매핑 안 된 상품:")
        for u in unmatched:
            print(f"  - {u['sku']}: {u['label']}")

    # 출력 파일명
    if not output_path:
        # 파일명에서 사업자명 추출 시도
        biz_name = '이더컴퍼니'
        import os
        base = os.path.basename(shipment_path)
        if '_' in base:
            parts = base.split('_')
            for p in parts:
                if '주식회사' in p or '컴퍼니' in p or '인테크' in p or '플로' in p or '코퍼레이션' in p:
                    biz_name = p.replace('주식회사', '').strip()
                    break
        output_path = f"{sqed_code}_{biz_name}_간편입고등록.xlsx"

    create_output(products, output_path)

    # 결과 출력
    total_qty = sum(p['qty'] for p in products)
    print(f"\n결과: {len(products)}개 상품, 총 수량 {total_qty}")
    for p in products:
        box_range = format_box_range(p['boxes'])
        print(f"  {box_range} | {p['prod_code']} | {p['prod_name']} | {p['qty']}")

    # 검증
    expected = int(float(orig_total_l)) if orig_total_l and float(orig_total_l) > 0 else int(float(orig_total_k or 0))
    if total_qty == expected:
        print(f"\n✅ 수량 검증 통과 (총 {total_qty}개 = 원본 {expected}개)")
    else:
        print(f"\n⚠️ 수량 불일치! 변환: {total_qty}개, 원본: {expected}개")

    print(f"✅ 저장 완료: {output_path}")


if __name__ == '__main__':
    main()
